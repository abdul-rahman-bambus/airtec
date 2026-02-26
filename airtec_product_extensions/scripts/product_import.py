# -*- coding: utf-8 -*-
import csv
import logging
from pathlib import Path

import openpyxl
from odoo.exceptions import ValidationError

_logger = logging.getLogger(__name__)

# ------------------------------------------------------------------
# CONFIGURATION
# ------------------------------------------------------------------

BASE_PATH = Path("/home/bambus/Desktop/airtec")
FILE_PATH = BASE_PATH / "products_v1.xlsx"

LOG_PATH = BASE_PATH / "import_logs"
LOG_PATH.mkdir(exist_ok=True)

BATCH_SIZE = 1
MAX_RECORDS = 30   # set None later to remove limit

# ------------------------------------------------------------------
# LOG FILES
# ------------------------------------------------------------------

SUCCESS_LOG = LOG_PATH / "product_import_success.log"
FAIL_LOG = LOG_PATH / "product_import_failed.log"
SUMMARY_LOG = LOG_PATH / "product_import_summary.log"
FAIL_CSV = LOG_PATH / "product_import_failed_rows.csv"

# ------------------------------------------------------------------
# ROUTES (STANDARD ODOO XML IDS)
# ------------------------------------------------------------------

ROUTE_XML_MAP = {
    "buy": "purchase_stock.route_warehouse0_buy",
    "manufacture": "mrp.route_warehouse0_manufacture",
    "mto": "stock.route_warehouse0_mto",
}

# ------------------------------------------------------------------
# UOM MAP (NORMALIZED)
# ------------------------------------------------------------------

UOM_MAP = {
    "stück": "Units",
    "stueck": "Units",
    "units": "Units",
    "unit": "Units",
    "kg": "kg",
    "g": "g",
    "m": "m",
    "m²": "m²",
}

# ------------------------------------------------------------------
# LOGGING HELPERS
# ------------------------------------------------------------------

def _append_text(path, line):
    with open(path, "a", encoding="utf-8") as f:
        f.write(line + "\n")


def log_success(msg):
    _append_text(SUCCESS_LOG, msg)


def log_fail(msg):
    _append_text(FAIL_LOG, msg)


def init_fail_csv():
    if not FAIL_CSV.exists():
        with open(FAIL_CSV, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow([
                "batch_no",
                "row_no",
                "default_code",
                "product_name",
                "column",
                "value",
                "reason",
            ])


def log_fail_csv(batch_no, row_no, row, column, value, reason):
    with open(FAIL_CSV, "a", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow([
            batch_no,
            row_no,
            row.get("default_code"),
            row.get("name"),
            column,
            value,
            reason,
        ])

def to_bool(val):
    if val is True:
        return True
    if val is False:
        return False
    if val is None:
        return False
    if isinstance(val, str):
        return val.strip().lower() in ("1", "true", "yes", "y")
    if isinstance(val, (int, float)):
        return val == 1
    return False



# ------------------------------------------------------------------
# XLSX READER
# ------------------------------------------------------------------

def read_xlsx(file_path):
    wb = openpyxl.load_workbook(file_path, data_only=True)
    sheet = wb.active

    headers = [str(c).strip() if c else "" for c in next(sheet.iter_rows(max_row=1, values_only=True))]

    rows = []
    for r in sheet.iter_rows(min_row=2, values_only=True):
        rows.append(dict(zip(headers, r)))

    return rows

# ------------------------------------------------------------------
# RESOLVERS
# ------------------------------------------------------------------

def resolve_uom(env, value):
    if not value:
        raise ValidationError("UoM value missing")

    key = str(value).strip().lower()
    mapped = UOM_MAP.get(key)

    if not mapped:
        raise ValidationError(f"UoM '{value}' not mapped")

    uom = env["uom.uom"].search([("name", "=", mapped)], limit=1)
    print(uom, 'uommmmmmm')
    if not uom:
        raise ValidationError(f"Mapped UoM '{mapped}' not found in Odoo")

    return uom.id


def resolve_routes(env, row):
    route_ids = []

    route_val = (row.get("route_ids") or "").strip()
    dispo = (row.get("Dispositionsart") or "").strip()

    if route_val == "Fremdbeschaffung":
        route_ids.append(env.ref(ROUTE_XML_MAP["buy"]).id)

    if route_val == "Eigenfertigung":
        route_ids.append(env.ref(ROUTE_XML_MAP["manufacture"]).id)

    if dispo == "auftragsbezogen":
        route_ids.append(env.ref(ROUTE_XML_MAP["mto"]).id)

    return list(set(route_ids))

def resolve_product_tags(env, row):
    tag_ids = []

    # TAG-* boolean columns ONLY
    for key, val in row.items():
        if key.startswith("TAG-") and val:
            tag_name = key.replace("TAG-", "").strip()
            tag = env["product.tag"].search([("name", "=", tag_name)], limit=1)
            if not tag:
                tag = env["product.tag"].create({"name": tag_name})
            tag_ids.append(tag.id)

    # Key:Value style tags (EXPLICITLY EXCLUDING Suchwort)
    for key, val in row.items():
        if ":" in key and val:
            normalized_key = key.strip().lower()

            # 🚫 BLOCK SUCHWORT COMPLETELY
            if normalized_key.startswith("suchwort"):
                continue

            tag_name = f"{key.strip()}: {str(val).strip()}"
            tag = env["product.tag"].search([("name", "=", tag_name)], limit=1)
            if not tag:
                tag = env["product.tag"].create({"name": tag_name})
            tag_ids.append(tag.id)

    return list(set(tag_ids))




def resolve_tracking(row):
    if row.get("serial"):
        return "serial"
    if row.get("lot"):
        return "lot"
    return "none"

# ------------------------------------------------------------------
# NOTES + SUCHWORT
# ------------------------------------------------------------------

def build_description(row, existing_desc=""):
    desc = (existing_desc or "").strip()

    base_note = (row.get("description") or "").strip()
    if base_note:
        desc = base_note

    suchwort = (row.get("Suchwort") or "").strip()
    if suchwort:
        line = f"Suchwort: {suchwort}"
        if line not in desc:
            if desc:
                desc += "\n\n"
            desc += line

    return desc

# ------------------------------------------------------------------
# PREPARE PRODUCT VALS
# ------------------------------------------------------------------

def prepare_product_vals(env, row, product=None):
    vals = {}

    vals["name"] = row.get("name")
    vals["default_code"] = row.get("default_code")
    vals["sale_ok"] = to_bool(row.get("sale_ok"))
    vals["purchase_ok"] = True

    vals["barcode"] = row.get("barcode")
    vals["description_purchase"] = row.get("description_purchase")
    vals["description_sale"] = row.get("description_sale")

    vals["atlas_relevant"] = bool(row.get("atlas_relevant"))
    vals["intrastat_relevant"] = bool(row.get("intrastat_relevant"))

    vals["device_label_line_4"] = row.get("device_label_line_4")
    vals["control_panel_cable_length_mm"] = row.get("control_panel_cable_length_mm")
    vals["coupling_cable_length_mm"] = row.get("coupling_cable_length_mm")

    # UOM
    uom_id = resolve_uom(env, row.get("uom_id"))
    vals["uom_id"] = uom_id
    vals["uom_po_id"] = uom_id

    # Tracking
    vals["tracking"] = resolve_tracking(row)

    # Routes
    route_ids = resolve_routes(env, row)
    if route_ids:
        vals["route_ids"] = [(6, 0, route_ids)]

    # Product Tags
    tag_ids = resolve_product_tags(env, row)
    if tag_ids:
        vals["product_tag_ids"] = [(6, 0, tag_ids)]

    # Description + Suchwort
    existing_desc = product.description if product and product.description else ""
    vals["description"] = build_description(row, existing_desc)

    return vals



# ------------------------------------------------------------------
# BATCH PROCESSOR
# ------------------------------------------------------------------

def process_batch(env, batch_rows, batch_no, stats):
    log_success(f"[BATCH {batch_no}] START")

    for row_no, row in batch_rows:
        try:
            default_code = row.get("default_code")
            if not default_code:
                raise ValidationError("Missing default_code")

            # ✅ Skip non-sale products (decision belongs here)
            _logger.warning(
                "DEBUG sale_ok raw=%r normalized=%s SKU=%s",
                row.get("sale_ok"),
                sale_ok,
                row.get("default_code"),
            )
            if not to_bool(row.get("sale_ok")):
                log_success(
                    f"[BATCH {batch_no}] Row {row_no} SKU={default_code} | skipped (sale_ok = false)"
                )
                continue


            product = env["product.template"].search(
                [("default_code", "=", default_code)], limit=1
            )

            vals = prepare_product_vals(env, row, product)

            if not vals:
                raise ValidationError("No values prepared for product")

            if product:
                product.write(vals)
                stats["updated"] += 1
                log_success(f"[BATCH {batch_no}] SKU={default_code} | updated")
            else:
                env["product.template"].create(vals)
                stats["created"] += 1
                log_success(f"[BATCH {batch_no}] SKU={default_code} | created")

        except Exception as e:
            stats["failed"] += 1
            log_fail(
                f"[BATCH {batch_no}] Row {row_no} SKU={row.get('default_code')} | {e}"
            )
            log_fail_csv(
                batch_no,
                row_no,
                row,
                "N/A",
                "",
                str(e),
            )

    log_success(
        f"[BATCH {batch_no}] END | created={stats['created']} "
        f"updated={stats['updated']} failed={stats['failed']}"
    )



# ------------------------------------------------------------------
# ENTRY POINT
# ------------------------------------------------------------------

def import_products(env):
    init_fail_csv()

    rows = read_xlsx(FILE_PATH)
    total_rows = len(rows)

    if MAX_RECORDS:
        rows = rows[:MAX_RECORDS]

    stats = {"created": 0, "updated": 0, "failed": 0}
    batch = []
    batch_no = 1

    for idx, row in enumerate(rows, start=2):
        batch.append((idx, row))

        if len(batch) == BATCH_SIZE:
            process_batch(env, batch, batch_no, stats)
            env.cr.commit()
            batch.clear()
            batch_no += 1

    if batch:
        process_batch(env, batch, batch_no, stats)
        env.cr.commit()

    _append_text(
        SUMMARY_LOG,
        (
            f"Total rows in file : {total_rows}\n"
            f"Rows processed     : {len(rows)}\n"
            f"Batch size         : {BATCH_SIZE}\n"
            f"Created            : {stats['created']}\n"
            f"Updated            : {stats['updated']}\n"
            f"Failed             : {stats['failed']}\n"
            "----------------------------------------"
        ),
    )

    return stats
